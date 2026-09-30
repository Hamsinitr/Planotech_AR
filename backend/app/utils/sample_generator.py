import os
import sys
from pathlib import Path

# Add backend directory to sys.path
backend_dir = Path(__file__).resolve().parent.parent.parent
if str(backend_dir) not in sys.path:
    sys.path.insert(0, str(backend_dir))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.config import settings

# Style helper
NAVY_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
WHITE_BOLD = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
THIN_SIDE = Side(border_style="thin", color="CBD5E1")
CELL_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

def style_excel(ws, headers, rows):
    ws.append(headers)
    ws.row_dimensions[1].height = 26
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col_idx)
        c.fill = NAVY_FILL
        c.font = WHITE_BOLD
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = CELL_BORDER

    for r_idx, row in enumerate(rows, start=2):
        ws.append(row)
        ws.row_dimensions[r_idx].height = 20
        for c_idx in range(1, len(row) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.border = CELL_BORDER
            cell.font = Font(name="Segoe UI", size=9)
            cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20

# 38 Employees across 4 departments: IT, HR, Marketing, Operations
EMPLOYEES = [
    # IT Department (12)
    {"id": "EMP001", "name": "Hamsini T R", "dept": "IT", "in": "09:12:00", "out": "18:05:00", "att": "Present", "has_rep": True, "rep_name": "HAMSINI T.R", "quality": "Excellent", "task": "Architected Azure Gamified Coding engine"},
    {"id": "EMP002", "name": "Aarav Sharma", "dept": "IT", "in": "09:25:00", "out": "18:30:00", "att": "Present", "has_rep": True, "rep_name": "Aarav Sharma", "quality": "Good", "task": "Implemented WebSocket leaderboard sync"},
    {"id": "EMP003", "name": "Diya Patel", "dept": "IT", "in": "09:35:00", "out": "18:15:00", "att": "Present", "has_rep": True, "rep_name": "Diya Patel", "quality": "Excellent", "task": "Created UI test harness with Cypress"},
    {"id": "EMP004", "name": "Rohan Iyer", "dept": "IT", "in": "09:52:00", "out": "19:00:00", "att": "Present", "has_rep": True, "rep_name": "Rohan Iyer", "quality": "Good", "task": "Optimized database query indexes (Late punch)"},
    {"id": "EMP005", "name": "Ananya Nair", "dept": "IT", "in": "09:20:00", "out": "18:10:00", "att": "Present", "has_rep": False, "rep_name": None, "quality": None, "task": None}, # Present + No report -> Not Submitted
    {"id": "EMP006", "name": "Vikram Malhotra", "dept": "IT", "in": "", "out": "", "att": "Absent", "has_rep": False, "rep_name": None, "quality": None, "task": None}, # Absent + No report -> Absent
    {"id": "EMP007", "name": "Pooja Reddy", "dept": "IT", "in": "09:15:00", "out": "18:00:00", "att": "Present", "has_rep": True, "rep_name": "Pooja Reddy", "quality": "Average", "task": "Frontend bugfixes on mobile navbar"},
    {"id": "EMP008", "name": "Siddharth Verma", "dept": "IT", "in": "09:40:00", "out": "18:45:00", "att": "Present", "has_rep": True, "rep_name": "Siddharth Verma", "quality": "Good", "task": "API security rate limiter integration"},
    {"id": "EMP009", "name": "Kavya Deshmukh", "dept": "IT", "in": "09:55:00", "out": "18:50:00", "att": "Present", "has_rep": True, "rep_name": "Kavya Deshmukh", "quality": "Poor", "task": "Partial refactor of payment flow (Late punch)"},
    {"id": "EMP010", "name": "Aditya Joshi", "dept": "IT", "in": "09:18:00", "out": "18:20:00", "att": "Present", "has_rep": True, "rep_name": "Aditya Joshi", "quality": "Excellent", "task": "Dockerized microservices deployment"},
    {"id": "EMP011", "name": "Meera Pillai", "dept": "IT", "in": "09:42:00", "out": "18:15:00", "att": "Present", "has_rep": False, "rep_name": None, "quality": None, "task": None}, # Present + No report -> Not Submitted
    {"id": "EMP012", "name": "Nikhil Bose", "dept": "IT", "in": "", "out": "", "att": "Absent", "has_rep": True, "rep_name": "Nikhil Bose", "quality": "Average", "task": "Submitted draft notes while on leave"}, # Absent + accidental report!

    # HR Department (8)
    {"id": "EMP013", "name": "Sunita Menon", "dept": "HR", "in": "09:10:00", "out": "17:45:00", "att": "Present", "has_rep": True, "rep_name": "Sunita Menon", "quality": "Excellent", "task": "Conducted 4 technical interviews"},
    {"id": "EMP014", "name": "Rajesh Kulkarni", "dept": "HR", "in": "09:28:00", "out": "18:00:00", "att": "Present", "has_rep": True, "rep_name": "Rajesh Kulkarni", "quality": "Good", "task": "Processed monthly payroll tax filings"},
    {"id": "EMP015", "name": "Deepa Singhania", "dept": "HR", "in": "09:33:00", "out": "18:10:00", "att": "Present", "has_rep": True, "rep_name": "Deepa Singhania", "quality": "Good", "task": "Organized annual wellness workshop"},
    {"id": "EMP016", "name": "Arjun Sen", "dept": "HR", "in": "09:50:00", "out": "18:30:00", "att": "Present", "has_rep": True, "rep_name": "Arjun Sen", "quality": "Average", "task": "Updated employee handbook policies (Late)"},
    {"id": "EMP017", "name": "Priya Choudhury", "dept": "HR", "in": "", "out": "", "att": "Absent", "has_rep": False, "rep_name": None, "quality": None, "task": None}, # Absent
    {"id": "EMP018", "name": "Kiran Kumar", "dept": "HR", "in": "09:14:00", "out": "17:50:00", "att": "Present", "has_rep": True, "rep_name": "Kiran Kumar", "quality": "Good", "task": "Campus recruitment drive planning"},
    {"id": "EMP019", "name": "Neha Agarwal", "dept": "HR", "in": "09:22:00", "out": "18:05:00", "att": "Present", "has_rep": False, "rep_name": None, "quality": None, "task": None}, # Not submitted
    {"id": "EMP020", "name": "Gaurav Mehta", "dept": "HR", "in": "09:16:00", "out": "17:55:00", "att": "Present", "has_rep": True, "rep_name": "Gaurav Mehta", "quality": "Excellent", "task": "Employee engagement pulse survey analysis"},

    # Marketing Department (10)
    {"id": "EMP021", "name": "Tarun Kapoor", "dept": "Marketing", "in": "09:05:00", "out": "18:00:00", "att": "Present", "has_rep": True, "rep_name": "Tarun Kapoor", "quality": "Excellent", "task": "Launched Q3 lead generation campaign"},
    {"id": "EMP022", "name": "Sneha Rao", "dept": "Marketing", "in": "09:26:00", "out": "18:15:00", "att": "Present", "has_rep": True, "rep_name": "Sneha Rao", "quality": "Good", "task": "Designed social media marketing kit"},
    {"id": "EMP023", "name": "Harsh Vardhan", "dept": "Marketing", "in": "09:38:00", "out": "18:20:00", "att": "Present", "has_rep": True, "rep_name": "Harsh Vardhan", "quality": "Average", "task": "Wrote two SEO blog articles"},
    {"id": "EMP024", "name": "Ishita Roy", "dept": "Marketing", "in": "10:02:00", "out": "19:00:00", "att": "Present", "has_rep": True, "rep_name": "Ishita Roy", "quality": "Good", "task": "Reviewed influencer marketing contracts (Late)"},
    {"id": "EMP025", "name": "Manish Tiwari", "dept": "Marketing", "in": "", "out": "", "att": "Absent", "has_rep": False, "rep_name": None, "quality": None, "task": None}, # Absent
    {"id": "EMP026", "name": "Ritu Saxena", "dept": "Marketing", "in": "09:11:00", "out": "17:45:00", "att": "Present", "has_rep": True, "rep_name": "Ritu Saxena", "quality": "Excellent", "task": "Analyzed Google Ads ROI metrics"},
    {"id": "EMP027", "name": "Abhishek Dubey", "dept": "Marketing", "in": "09:44:00", "out": "18:15:00", "att": "Present", "has_rep": True, "rep_name": "Abhishek Dubey", "quality": "Average", "task": "Email newsletter layout draft"},
    {"id": "EMP028", "name": "Pallavi Nambiar", "dept": "Marketing", "in": "09:29:00", "out": "18:00:00", "att": "Present", "has_rep": False, "rep_name": None, "quality": None, "task": None}, # Not submitted
    {"id": "EMP029", "name": "Varun Sethi", "dept": "Marketing", "in": "09:19:00", "out": "18:05:00", "att": "Present", "has_rep": True, "rep_name": "Varun Sethi", "quality": "Good", "task": "Customer webinar presentation deck"},
    {"id": "EMP030", "name": "Divya Swaminathan", "dept": "Marketing", "in": "09:58:00", "out": "19:10:00", "att": "Present", "has_rep": True, "rep_name": "Divya Swaminathan", "quality": "Poor", "task": "Competitor pricing research (Late)"},

    # Operations Department (8)
    {"id": "EMP031", "name": "Karthik Ranganathan", "dept": "Operations", "in": "09:08:00", "out": "18:00:00", "att": "Present", "has_rep": True, "rep_name": "Karthik Ranganathan", "quality": "Excellent", "task": "Vendor contract negotiation completed"},
    {"id": "EMP032", "name": "Shweta Pandey", "dept": "Operations", "in": "09:24:00", "out": "18:15:00", "att": "Present", "has_rep": True, "rep_name": "Shweta Pandey", "quality": "Good", "task": "Facility inventory check audit"},
    {"id": "EMP033", "name": "Naveen Chawla", "dept": "Operations", "in": "09:36:00", "out": "18:25:00", "att": "Present", "has_rep": True, "rep_name": "Naveen Chawla", "quality": "Good", "task": "Logistics dispatch route optimization"},
    {"id": "EMP034", "name": "Bhavna Jain", "dept": "Operations", "in": "10:15:00", "out": "19:30:00", "att": "Present", "has_rep": True, "rep_name": "Bhavna Jain", "quality": "Average", "task": "IT hardware asset tracking (Late)"},
    {"id": "EMP035", "name": "Suresh Bhat", "dept": "Operations", "in": "09:17:00", "out": "17:50:00", "att": "Present", "has_rep": True, "rep_name": "Suresh Bhat", "quality": "Good", "task": "Quarterly procurement budget planning"},
    {"id": "EMP036", "name": "Tanvi Goswami", "dept": "Operations", "in": "09:21:00", "out": "18:05:00", "att": "Present", "has_rep": True, "rep_name": "Tanvi Goswami", "quality": "Excellent", "task": "ISO compliance documentation checklist"},
    {"id": "EMP037", "name": "Rakesh Chauhan", "dept": "Operations", "in": "09:41:00", "out": "18:30:00", "att": "Present", "has_rep": True, "rep_name": "Rakesh Chauhan", "quality": "Average", "task": "Client onboarding facility setup"},
    {"id": "EMP038", "name": "Lavanya Krishnan", "dept": "Operations", "in": "09:27:00", "out": "18:10:00", "att": "Present", "has_rep": True, "rep_name": "Lavanya Krishnan", "quality": "Good", "task": "Office safety drill coordination"}
]

def generate_sample_excel_files(target_dir: str, date_str: str = "2026-09-15"):
    """Generates realistic sample Attendance and Report Excel workbooks."""
    os.makedirs(target_dir, exist_ok=True)

    # 1. Attendance Excel
    att_wb = openpyxl.Workbook()
    att_ws = att_wb.active
    att_ws.title = "Attendance"

    att_headers = ["Employee ID", "Employee Name", "Department", "Punch In Time", "Punch Out Time", "Attendance Status"]
    att_rows = []
    for emp in EMPLOYEES:
        att_rows.append([
            emp["id"],
            emp["name"],
            emp["dept"],
            emp["in"],
            emp["out"],
            emp["att"]
        ])
    style_excel(att_ws, att_headers, att_rows)
    att_path = os.path.join(target_dir, f"Attendance_{date_str}.xlsx")
    att_wb.save(att_path)

    # 2. Daily Report Excel
    rep_wb = openpyxl.Workbook()
    rep_ws = rep_wb.active
    rep_ws.title = "Daily Work Reports"

    rep_headers = ["Employee ID", "Employee Name", "Department", "Submission Status", "Report Quality", "Task Summary"]
    rep_rows = []
    for emp in EMPLOYEES:
        if emp["has_rep"]:
            rep_rows.append([
                emp["id"],
                emp["rep_name"], # May have slight formatting difference like 'HAMSINI T.R'
                emp["dept"],
                "Submitted",
                emp["quality"],
                emp["task"]
            ])
    style_excel(rep_ws, rep_headers, rep_rows)
    rep_path = os.path.join(target_dir, f"Daily_Report_{date_str}.xlsx")
    rep_wb.save(rep_path)

    print(f"Sample files generated:\n- {att_path}\n- {rep_path}")
    return att_path, rep_path

if __name__ == "__main__":
    generate_sample_excel_files(settings.UPLOAD_DIR)
