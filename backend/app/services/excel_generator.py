import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Styling constants
NAVY_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
WHITE_BOLD_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

ROW_FONT = Font(name="Segoe UI", size=9)
ROW_BOLD_FONT = Font(name="Segoe UI", size=9, bold=True)
TOTAL_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

THIN_BORDER_SIDE = Side(border_style="thin", color="CBD5E1")
CELL_BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)

# Status color fills
STATUS_FILLS = {
    "present": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),       # Green
    "submitted": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
    "excellent": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
    "top performer": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
    "on time": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),

    "late": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),          # Amber
    "grace (9:30-9:45)": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
    "average": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
    "good": PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid"),          # Soft blue

    "not submitted": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),  # Red
    "poor": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
    "needs attention": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),

    "absent": PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),        # Slate
}

def style_sheet(ws, headers: list[str], data_rows: list[list], is_total_row_present: bool = True):
    """Applies standardized corporate styling, auto-column widths, and freeze panes."""
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"

    # Header Row
    ws.append(headers)
    ws.row_dimensions[1].height = 28
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = NAVY_FILL
        cell.font = WHITE_BOLD_FONT
        cell.alignment = HEADER_ALIGNMENT
        cell.border = CELL_BORDER

    # Data Rows
    for row_idx, row_data in enumerate(data_rows, start=2):
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 20
        is_total = is_total_row_present and (row_idx == len(data_rows) + 1)
        zebra = ZEBRA_FILL if (row_idx % 2 == 0 and not is_total) else WHITE_FILL

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = CELL_BORDER
            val_str = str(cell.value).strip().lower() if cell.value is not None else ""

            if is_total:
                cell.fill = TOTAL_FILL
                cell.font = ROW_BOLD_FONT
            else:
                cell.font = ROW_FONT
                # Check status fill
                if val_str in STATUS_FILLS:
                    cell.fill = STATUS_FILLS[val_str]
                else:
                    cell.fill = zebra

            # Number formatting
            if isinstance(cell.value, float):
                cell.number_format = "0.0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif isinstance(cell.value, int):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

def generate_processed_excel(
    date_str: str,
    raw_attendance_rows: list[dict],
    att_summary: dict,
    rep_summary: dict,
    attendance_details: list[dict],
    report_details: list[dict],
    productivity_data: dict,
    output_path: str
) -> str:
    """Generates the master 9-Sheet processed Excel file."""
    wb = Workbook()
    wb.remove(wb.active)  # Remove default blank sheet

    # 1. Sheet 1: Original Attendance Data
    ws1 = wb.create_sheet(title="Original Attendance Data")
    headers1 = ["Employee ID", "Employee Name", "Department", "Punch In Time", "Punch Out Time", "Status"]
    rows1 = [
        [r.get("employee_id") or "N/A", r.get("employee_name"), r.get("department"), r.get("punch_in_time") or "-", r.get("punch_out_time") or "-", r.get("attendance_status")]
        for r in raw_attendance_rows
    ]
    style_sheet(ws1, headers1, rows1, is_total_row_present=False)

    # 2. Sheet 2: Attendance Department Summary
    ws2 = wb.create_sheet(title="Att Dept Summary")
    headers2 = ["Department", "Total Individuals", "Before 9:30 AM", "9:30 AM–9:45 AM", "After 9:45 AM"]
    rows2 = [
        [r["department"], r["total_individuals"], r["before_930"], r["between_930_945"], r["after_945"]]
        for r in att_summary.get("summary_1", [])
    ]
    tot1 = att_summary.get("summary_1_total", {})
    rows2.append(["Total", tot1.get("total_individuals", 0), tot1.get("before_930", 0), tot1.get("between_930_945", 0), tot1.get("after_945", 0)])
    style_sheet(ws2, headers2, rows2, is_total_row_present=True)

    # 3. Sheet 3: Attendance Status Summary
    ws3 = wb.create_sheet(title="Att Status Summary")
    headers3 = ["Department", "Total Individuals", "Not Punched In", "Present", "Late Punch-In"]
    rows3 = [
        [r["department"], r["total_individuals"], r["not_punched_in"], r["present"], r["late_punch_in"]]
        for r in att_summary.get("summary_2", [])
    ]
    tot2 = att_summary.get("summary_2_total", {})
    rows3.append(["Total", tot2.get("total_individuals", 0), tot2.get("not_punched_in", 0), tot2.get("present", 0), tot2.get("late_punch_in", 0)])
    style_sheet(ws3, headers3, rows3, is_total_row_present=True)

    # 4. Sheet 4: Employee Attendance Details
    ws4 = wb.create_sheet(title="Emp Attendance Details")
    headers4 = ["Employee Name", "Employee ID", "Department", "Date", "Punch In", "Punch Out", "Attendance Status", "Late Status"]
    rows4 = [
        [r.get("employee_name"), r.get("employee_id") or "N/A", r.get("department"), date_str, r.get("punch_in_time") or "-", r.get("punch_out_time") or "-", r.get("attendance_status"), r.get("late_status")]
        for r in attendance_details
    ]
    style_sheet(ws4, headers4, rows4, is_total_row_present=False)

    # 5. Sheet 5: Report Submission Department Summary
    ws5 = wb.create_sheet(title="Report Dept Summary")
    headers5 = ["Department", "Total Individuals", "Reports Submitted", "Reports Not Submitted", "Absent"]
    rows5 = [
        [r["department"], r["total_individuals"], r["reports_submitted"], r["reports_not_submitted"], r["absent"]]
        for r in rep_summary.get("summary", [])
    ]
    tot_rep = rep_summary.get("summary_total", {})
    rows5.append(["Total", tot_rep.get("total_individuals", 0), tot_rep.get("reports_submitted", 0), tot_rep.get("reports_not_submitted", 0), tot_rep.get("absent", 0)])
    style_sheet(ws5, headers5, rows5, is_total_row_present=True)

    # 6. Sheet 6: Employee Report Status
    ws6 = wb.create_sheet(title="Emp Report Status")
    headers6 = ["Employee Name", "Employee ID", "Department", "Date", "Attendance Status", "Report Status", "Report Quality", "Notes"]
    rows6 = [
        [r.get("employee_name"), r.get("employee_id") or "N/A", r.get("department"), date_str, r.get("attendance_status"), r.get("report_status"), r.get("quality_rating") or "-", r.get("notes") or "-"]
        for r in report_details
    ]
    style_sheet(ws6, headers6, rows6, is_total_row_present=False)

    # 7. Sheet 7: Report Quality Analysis
    ws7 = wb.create_sheet(title="Report Quality Analysis")
    headers7 = ["Department", "Total Reports", "Excellent", "Good", "Average", "Poor", "Average Quality Score"]
    rows7 = [
        [r["department"], r["total_reports"], r["excellent"], r["good"], r["average"], r["poor"], r["avg_quality_score"]]
        for r in rep_summary.get("quality_summary", [])
    ]
    tot_qual = rep_summary.get("quality_total", {})
    rows7.append(["Total", tot_qual.get("total_reports", 0), tot_qual.get("excellent", 0), tot_qual.get("good", 0), tot_qual.get("average", 0), tot_qual.get("poor", 0), tot_qual.get("avg_quality_score", 0.0)])
    style_sheet(ws7, headers7, rows7, is_total_row_present=True)

    # 8. Sheet 8: Employee Productivity
    ws8 = wb.create_sheet(title="Employee Productivity")
    headers8 = ["Rank", "Employee Name", "Employee ID", "Department", "Report Submission Score", "Quality Score", "Attendance Score", "Productivity Score", "Category"]
    rows8 = [
        [r.get("rank"), r.get("employee_name"), r.get("employee_id") or "N/A", r.get("department"), r.get("submission_score", 0.0), r.get("quality_score", 0.0), r.get("attendance_score", 0.0), r.get("productivity_score", 0.0), r.get("category")]
        for r in productivity_data.get("all_rankings", [])
    ]
    style_sheet(ws8, headers8, rows8, is_total_row_present=False)

    # 9. Sheet 9: Department Productivity
    ws9 = wb.create_sheet(title="Department Productivity")
    headers9 = ["Department", "Employees", "Present", "Reports Submitted", "Reports Not Submitted", "Avg Quality", "Productivity Score"]
    rows9 = [
        [r["department"], r["employees"], r["present"], r["reports_submitted"], r["reports_not_submitted"], r["avg_quality"], r["productivity_score"]]
        for r in productivity_data.get("department_summary", [])
    ]
    style_sheet(ws9, headers9, rows9, is_total_row_present=False)

    # Save
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path
